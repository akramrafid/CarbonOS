import os
import hashlib
import json
import re
from typing import List, Dict, Any, Optional, NamedTuple
from sqlalchemy.orm import Session
from .config import (
    CHROMA_PERSIST_DIR,
    CHUNK_SIZE,
    CHUNK_OVERLAP,
    EMBEDDING_MODEL,
    CORPUS_REGULATORY,
    CORPUS_EMISSION_FACTORS,
    CORPUS_TENANT_DOCS,
    VALID_CORPORA
)
from .models import ESGDocument, ESGChunk

# Vector store client singleton
_chroma_client = None
_embedding_model = None

def get_chroma_client():
    global _chroma_client
    if _chroma_client is None:
        try:
            import chromadb
            os.makedirs(CHROMA_PERSIST_DIR, exist_ok=True)
            _chroma_client = chromadb.PersistentClient(path=CHROMA_PERSIST_DIR)
        except Exception as e:
            print(f"[ESG RAG] ChromaDB init error ({e}), using in-memory/file fallback")
            _chroma_client = None
    return _chroma_client

def get_embedding_model():
    global _embedding_model
    if _embedding_model is None:
        try:
            from sentence_transformers import SentenceTransformer
            _embedding_model = SentenceTransformer(EMBEDDING_MODEL)
        except Exception as e:
            print(f"[ESG RAG] SentenceTransformer not loaded ({e}), using fast hash-vector fallback")
            _embedding_model = None
    return _embedding_model

def compute_embedding(text: str) -> List[float]:
    """Generates 384-dimensional dense embedding vector."""
    model = get_embedding_model()
    if model is not None:
        try:
            emb = model.encode(text, convert_to_numpy=True)
            return emb.tolist()
        except Exception as e:
            print(f"[ESG RAG] Error encoding with model: {e}")
            
    # Deterministic semantic hash fallback (384-dim normalized)
    import numpy as np
    words = re.findall(r'\w+', text.lower())
    vec = np.zeros(384, dtype=np.float32)
    for word in words:
        h = int(hashlib.md5(word.encode('utf-8')).hexdigest(), 16)
        idx = h % 384
        sign = 1.0 if ((h >> 9) & 1) else -1.0
        vec[idx] += sign
    norm = np.linalg.norm(vec)
    if norm > 1e-6:
        vec = vec / norm
    return vec.tolist()

def extract_text_from_file(file_path: str, filename: str) -> List[Dict[str, Any]]:
    """
    Extracts text and page numbers from PDF, XLSX, CSV, TXT, or MD files.
    Returns list of {"page": int, "text": str}.
    """
    pages_data = []
    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext == ".pdf":
            # Scanned PDFs yield no text: callers must report that, not parse raw bytes.
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for i, p in enumerate(pdf.pages):
                    txt = p.extract_text() or ""
                    if txt.strip():
                        pages_data.append({"page": i + 1, "text": txt})

        elif ext in [".xlsx", ".xls"]:
            import pandas as pd
            excel = pd.ExcelFile(file_path)
            for i, sheet_name in enumerate(excel.sheet_names):
                df = excel.parse(sheet_name)
                txt = f"Sheet: {sheet_name}\n" + df.to_string()
                pages_data.append({"page": i + 1, "text": txt})

        elif ext == ".csv":
            import pandas as pd
            df = pd.read_csv(file_path)
            pages_data.append({"page": 1, "text": df.to_string()})

        elif ext == ".docx":
            import docx
            doc = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            pages_data.append({"page": 1, "text": full_text})
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                pages_data.append({"page": 1, "text": f.read()})

    except Exception as e:
        # Never index the error text as if it were document content.
        print(f"[ESG RAG] Text extraction failed for {filename}: {type(e).__name__}")
        return []

    return [p for p in pages_data if p["text"].strip()]

def chunk_text(pages: List[Dict[str, Any]], chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[Dict[str, Any]]:
    """Splits document pages into overlapping semantic text chunks."""
    chunks = []
    chunk_idx = 0

    for item in pages:
        page_num = item["page"]
        text = item["text"]
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            # Try to break on sentence or newline if possible
            if end < text_len:
                last_newline = text.rfind('\n', start, end)
                last_period = text.rfind('. ', start, end)
                split_point = max(last_newline, last_period)
                if split_point > start + (chunk_size // 2):
                    end = split_point + 1

            chunk_str = text[start:end].strip()
            if chunk_str:
                chunks.append({
                    "chunk_index": chunk_idx,
                    "content": chunk_str,
                    "page_number": page_num,
                    "token_count": len(chunk_str.split())
                })
                chunk_idx += 1

            start = end - overlap if (end - overlap) > start else end

    return chunks

ALLOWED_EXTENSIONS = {".pdf", ".csv", ".xlsx", ".xls", ".txt", ".md", ".docx"}
MAX_UPLOAD_BYTES = int(os.getenv("ESG_MAX_UPLOAD_BYTES", str(20 * 1024 * 1024)))


class IngestRejected(ValueError):
    """Upload refused before anything is stored."""


class IngestResult(NamedTuple):
    document: ESGDocument
    pages: List[Dict[str, Any]]
    duplicate: bool


def safe_filename(filename: Optional[str]) -> str:
    """Display-safe basename. Never used to build a storage path."""
    base = re.split(r"[\\/]", filename or "")[-1]
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base).strip(" .")[:120]
    ext = os.path.splitext(base)[1].lower()
    if not base or ext not in ALLOWED_EXTENSIONS:
        raise IngestRejected(f"Unsupported file type. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    return base


def _storage_path(tenant_id: str, checksum: str, ext: str) -> str:
    """Server-generated path: uploads/<tenant>/<sha256><ext>, contained in the upload root."""
    root = os.path.realpath(os.path.join(CHROMA_PERSIST_DIR, "uploads"))
    tenant_dir = re.sub(r"[^A-Za-z0-9_-]", "_", tenant_id)
    path = os.path.realpath(os.path.join(root, tenant_dir, checksum + ext))
    if not path.startswith(root + os.sep):
        raise IngestRejected("Invalid storage location")
    return path


def ingest_document(
    db: Session,
    file_bytes: bytes,
    filename: str,
    corpus: str = CORPUS_TENANT_DOCS,
    tenant_id: str = "default_tenant",
    meta_info: Optional[Dict[str, Any]] = None
) -> IngestResult:
    """
    Ingests document into the audit tables and vector store for one tenant.
    The caller is responsible for deriving tenant_id from an authenticated principal.
    Re-uploading identical bytes to the same tenant/corpus returns the existing document.
    """
    if corpus not in VALID_CORPORA:
        raise IngestRejected("Unknown corpus")
    filename = safe_filename(filename)
    if not file_bytes:
        raise IngestRejected("Empty file")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise IngestRejected(f"File exceeds {MAX_UPLOAD_BYTES} bytes")

    ext = os.path.splitext(filename)[1].lower()
    checksum = hashlib.sha256(file_bytes).hexdigest()
    path = _storage_path(tenant_id, checksum, ext)

    existing = db.query(ESGDocument).filter(
        ESGDocument.tenant_id == tenant_id,
        ESGDocument.corpus == corpus,
        ESGDocument.checksum == checksum,
    ).first()
    if existing is not None and os.path.exists(path):
        return IngestResult(existing, extract_text_from_file(path, filename), True)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f:
        f.write(file_bytes)

    pages = extract_text_from_file(path, filename)
    chunk_data = chunk_text(pages)

    # Create document record
    doc_rec = ESGDocument(
        tenant_id=tenant_id,
        corpus=corpus,
        filename=filename,
        file_type=ext.replace('.', ''),
        file_size=len(file_bytes),
        checksum=checksum,
        chunk_count=len(chunk_data),
        meta_info=meta_info or {}
    )
    db.add(doc_rec)
    db.flush()

    # ChromaDB indexing
    chroma = get_chroma_client()
    collection = None
    if chroma is not None:
        try:
            collection = chroma.get_or_create_collection(
                name=f"esg_{corpus}",
                metadata={"hnsw:space": "cosine"}
            )
        except Exception as ce:
            print(f"[ESG RAG] Chroma collection fetch error: {ce}")

    # Fallback file storage for vectors
    fallback_store_file = os.path.join(CHROMA_PERSIST_DIR, f"fallback_{corpus}.json")

    ids = []
    embeddings = []
    metadatas = []
    documents = []

    for c in chunk_data:
        chunk_rec = ESGChunk(
            document_id=doc_rec.id,
            chunk_index=c["chunk_index"],
            content=c["content"],
            token_count=c["token_count"],
            page_number=c["page_number"],
            vector_id=f"{doc_rec.id}_{c['chunk_index']}"
        )
        db.add(chunk_rec)

        vec = compute_embedding(c["content"])
        chunk_id = f"{doc_rec.id}_{c['chunk_index']}"
        ids.append(chunk_id)
        embeddings.append(vec)
        documents.append(c["content"])
        metadatas.append({
            "doc_id": doc_rec.id,
            "tenant_id": tenant_id,
            "filename": filename,
            "corpus": corpus,
            "page": c["page_number"]
        })

    if collection is not None and ids:
        try:
            collection.add(
                ids=ids,
                embeddings=embeddings,
                metadatas=metadatas,
                documents=documents
            )
        except Exception as e:
            print(f"[ESG RAG] Error saving to ChromaDB: {e}")

    # Always write to fallback store as secondary persistence
    try:
        existing = []
        if os.path.exists(fallback_store_file):
            with open(fallback_store_file, "r", encoding="utf-8") as f:
                existing = json.load(f)
        for i in range(len(ids)):
            existing.append({
                "id": ids[i],
                "vector": embeddings[i],
                "metadata": metadatas[i],
                "document": documents[i]
            })
        with open(fallback_store_file, "w", encoding="utf-8") as f:
            json.dump(existing, f)
    except Exception as fe:
        print(f"[ESG RAG] Fallback store error: {fe}")

    db.commit()
    db.refresh(doc_rec)
    return IngestResult(doc_rec, pages, False)
